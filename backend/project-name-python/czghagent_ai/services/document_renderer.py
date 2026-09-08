# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-08-03
import io
import zipfile
from datetime import UTC, datetime
from typing import cast
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from czghagent_ai.document_models import DocumentQaResult, DocumentRenderRequest
from czghagent_ai.services.document_html import HtmlBlock, parse_tender_html
from czghagent_ai.services.document_qa import DocumentQualityAssurance
from czghagent_ai.services.document_styles import (
    INK,
    MUTED,
    LayoutProfile,
    add_body_paragraph,
    add_data_table,
    add_list,
    configure_document,
    set_run_font,
)
from czghagent_ai.services.document_toc import (
    add_heading_bookmark,
    add_reference_toc,
    heading_bookmark_name,
)
from czghagent_ai.services.outline_numbering import build_outline_labels, order_outline
from czghagent_ai.services.table_format import (
    build_fallback_table_title,
    build_numbered_table_title,
)


class DocumentRenderer:
    def __init__(self, quality_assurance: DocumentQualityAssurance | None = None) -> None:
        self._quality_assurance = quality_assurance or DocumentQualityAssurance()

    def render(self, request: DocumentRenderRequest) -> tuple[bytes, DocumentQaResult]:
        if any(not chapter.content.strip() for chapter in request.chapters):
            raise ValueError("存在空白正文章节，不能开始正式排版")
        sanitized, qa = self._render_profile(request, "standard")
        applied_profile: LayoutProfile = "standard"
        for retry_profile in _page_retry_profiles(qa, request.target_pages):
            sanitized, qa = self._render_profile(request, retry_profile)
            applied_profile = retry_profile
            if qa.status == "PASSED" or qa.blank_pages or qa.forbidden_hits:
                break
        if applied_profile != "standard":
            qa = qa.model_copy(update={
                "summary": f"{qa.summary}；已自动采用{_profile_label(applied_profile)}排版"
            })
        return sanitized, qa

    def _render_profile(
        self, request: DocumentRenderRequest, profile: LayoutProfile
    ) -> tuple[bytes, DocumentQaResult]:
        draft = self._build_document(request, None, profile)
        page_numbers = self._quality_assurance.resolve_outline_pages(
            draft, order_outline(request.outline)
        )
        sanitized = self._build_document(request, page_numbers, profile)
        qa = self._quality_assurance.inspect(
            sanitized, request.target_pages, request.forbidden_terms
        )
        return sanitized, qa

    def _build_document(
        self, request: DocumentRenderRequest, page_numbers: list[int] | None,
        profile: LayoutProfile = "standard",
    ) -> bytes:
        document = Document()
        configure_document(document, profile)
        self._set_metadata(document, request)
        self._add_cover(document, request)
        document.add_page_break()  # type: ignore[no-untyped-call]
        add_reference_toc(document, request.outline, page_numbers)
        document.add_page_break()  # type: ignore[no-untyped-call]
        self._add_content(document, request, profile)
        buffer = io.BytesIO()
        document.save(buffer)
        return scrub_private_metadata(buffer.getvalue())

    def _set_metadata(self, document: DocumentObject, request: DocumentRenderRequest) -> None:
        properties = document.core_properties
        properties.title = request.title
        properties.subject = "技术投标文件"
        properties.author = ""
        properties.last_modified_by = ""
        properties.keywords = ""
        properties.comments = ""
        properties.category = ""
        properties.created = datetime(2000, 1, 1, tzinfo=UTC)
        properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
        properties.revision = 1

    def _add_cover(self, document: DocumentObject, request: DocumentRenderRequest) -> None:
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_before = Pt(105)
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(22)
        title.paragraph_format.line_spacing = 1.25
        title_run = title.add_run(_balance_cover_title(request.title))
        set_run_font(title_run, "黑体", 26, INK, bold=True)
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(12)
        subtitle_run = subtitle.add_run(
            "技术投标文件（暗标）" if request.bidding_mode == "BLIND" else "技术投标文件（明标）"
        )
        set_run_font(subtitle_run, "黑体", 16, MUTED)
        if request.bidding_mode == "BLIND":
            note = document.add_paragraph()
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            note.paragraph_format.space_before = Pt(110)
            note_run = note.add_run("技术部分")
            set_run_font(note_run, "宋体", 12, MUTED)

    def _add_content(
        self, document: DocumentObject, request: DocumentRenderRequest,
        profile: LayoutProfile,
    ) -> None:
        chapters = {chapter.outline_node_id: chapter for chapter in request.chapters}
        labels = build_outline_labels(request.outline)
        chapter_number = 0
        table_number = 0
        for bookmark_id, node in enumerate(order_outline(request.outline), 1):
            if node.level == 1:
                chapter_number += 1
                table_number = 0
            heading = document.add_heading(labels[node.id], level=node.level)
            for run in heading.runs:
                set_run_font(
                    run, "黑体", {1: 18, 2: 16, 3: 14}[node.level], INK, bold=True
                )
            add_heading_bookmark(
                heading, heading_bookmark_name(node.id), bookmark_id
            )
            chapter = chapters.get(node.id)
            if chapter is None:
                continue
            chapter_table_index = 0
            for block in parse_tender_html(chapter.content):
                if block.kind == "table":
                    table_number += 1
                    chapter_table_index += 1
                    fallback_title = build_fallback_table_title(
                        block.header, node.title, chapter_table_index
                    )
                    block.caption = build_numbered_table_title(
                        block.caption, chapter_number or 1, table_number, fallback_title
                    )
                self._add_block(document, block, profile)

    def _add_block(
        self, document: DocumentObject, block: HtmlBlock, profile: LayoutProfile
    ) -> None:
        if block.kind == "paragraph":
            add_body_paragraph(document, block.text, profile)
        elif block.kind == "heading":
            add_body_paragraph(document, block.text, profile)
        elif block.kind in {"bullet_list", "ordered_list"}:
            add_list(document, block.items, block.kind == "ordered_list", profile)
        elif block.kind == "table":
            add_data_table(
                document, block.header, block.rows, block.caption, block.note, profile
            )


def _page_retry_profiles(
    qa: DocumentQaResult, target_pages: int
) -> tuple[LayoutProfile, ...]:
    if qa.status != "FAILED" or qa.actual_pages is None:
        return ()
    if qa.blank_pages or qa.forbidden_hits:
        return ()
    if qa.actual_pages > target_pages:
        return ("compact", "condensed", "dense")
    if qa.actual_pages < target_pages:
        return ("spacious",)
    return ()


def _profile_label(profile: LayoutProfile) -> str:
    return {
        "standard": "标准", "compact": "紧凑", "condensed": "适度压缩",
        "dense": "密排", "spacious": "宽松",
    }[profile]


def scrub_private_metadata(content: bytes) -> bytes:
    source = io.BytesIO(content)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as input_zip, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as output_zip:
        for item in input_zip.infolist():
            if item.filename == "docProps/custom.xml":
                continue
            data = input_zip.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                data = _strip_revision_ids(data)
            output_zip.writestr(item, data)
    return output.getvalue()


def _strip_revision_ids(data: bytes) -> bytes:
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError:
        return data
    changed = False
    for element in root.iter():
        for attribute in list(element.attrib):
            local_name = attribute.rsplit("}", 1)[-1]
            if local_name.startswith("rsid"):
                del element.attrib[attribute]
                changed = True
    if not changed:
        return data
    ElementTree.register_namespace("w", "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
    ElementTree.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    return cast(bytes, ElementTree.tostring(root, encoding="utf-8", xml_declaration=True))


def _balance_cover_title(title: str) -> str:
    if "\n" in title or len(title) <= 16:
        return title
    midpoint = len(title) // 2
    boundaries = _semantic_boundaries(title)
    safe_boundaries = [value for value in boundaries if _is_safe_title_break(title, value)]
    candidates = safe_boundaries or [
        value for value in range(1, len(title)) if _is_safe_title_break(title, value)
    ]
    split_at = min(candidates, key=lambda value: abs(value - midpoint)) if candidates else midpoint
    return f"{title[:split_at]}\n{title[split_at:]}"


def _semantic_boundaries(title: str) -> list[int]:
    terms = ("生产运营", "总体", "技术", "平台", "系统", "方案", "建设", "服务", "页", "暗标", "项目")
    lower_bound = round(len(title) * 0.35)
    upper_bound = round(len(title) * 0.65)
    boundaries = [
        index + len(term)
        for term in terms
        for index in [title.find(term)]
        if index >= 0 and lower_bound <= index + len(term) <= upper_bound
    ]
    return boundaries


def _is_safe_title_break(title: str, index: int) -> bool:
    if index <= 0 or index >= len(title):
        return False
    left = title[index - 1]
    right = title[index]
    if left.isascii() and left.isalnum() and right.isascii() and right.isalnum():
        return False
    if right in "-—_:：/）)]}" or left in "-—_:：/（([{":
        return False
    if left.isascii() and left.isdigit() and right in "页年月日号":
        return False
    return True
