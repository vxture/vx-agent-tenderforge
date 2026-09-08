# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-07-29
import csv
import io
import os
import shutil
import subprocess
import tempfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import fitz  # type: ignore[import-untyped]
import pytesseract  # type: ignore[import-untyped]
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image


class UnsupportedDocumentError(ValueError):
    """文件格式不受支持或正文无法读取。"""


@dataclass(frozen=True)
class TextSegment:
    locator_type: str
    locator: str
    text: str
    ocr_confidence: float | None = None


class FileReaders:
    def __init__(self) -> None:
        self._readers: dict[str, Callable[[bytes], list[TextSegment]]] = {
            ".docx": self._read_docx,
            ".pdf": self._read_pdf,
            ".xlsx": self._read_xlsx,
            ".xlsm": self._read_xlsx,
            ".csv": self._read_csv,
            ".txt": self._read_text,
            ".md": self._read_text,
        }

    def read(self, file_name: str, content: bytes) -> list[TextSegment]:
        """
        将支持的文件转换为带来源定位的文本片段。

        Preconditions:
            - file_name 包含真实扩展名，content 为完整文件字节。
        Side Effects:
            - 旧版 DOC 解析时短暂创建临时文件并调用系统转换器。
        Error Semantics:
            - 不支持格式或无可读正文时抛出 UnsupportedDocumentError。
        """
        suffix = Path(file_name).suffix.lower()
        if suffix == ".doc":
            segments = self._read_legacy_doc(content)
        else:
            reader = self._readers.get(suffix)
            if reader is None:
                raise UnsupportedDocumentError(f"暂不支持 {suffix or '无扩展名'} 文件")
            segments = reader(content)
        cleaned = [segment for segment in segments if segment.text.strip()]
        if not cleaned:
            raise UnsupportedDocumentError("文件中没有可读取的文字内容")
        return cleaned

    def _read_docx(self, content: bytes) -> list[TextSegment]:
        document = Document(io.BytesIO(content))
        result: list[TextSegment] = []
        paragraph_index = 0
        table_index = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_index += 1
                text = block.text.strip()
                if not text:
                    continue
                style_name = block.style.name if block.style is not None else ""
                locator_type = "HEADING" if style_name.lower().startswith("heading") else "PARAGRAPH"
                result.append(TextSegment(locator_type, f"段落 {paragraph_index}", text))
                continue
            if not isinstance(block, Table):
                continue
            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                result.append(TextSegment(
                    "SHEET_ROW", f"表 {table_index} 行 {row_index}", " | ".join(values)
                ))
        return result

    def _read_pdf(self, content: bytes) -> list[TextSegment]:
        document = fitz.open(stream=content, filetype="pdf")
        result: list[TextSegment] = []
        try:
            for index, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if len(text) >= 40:
                    result.append(TextSegment("PAGE", f"第 {index} 页", text))
                    continue
                ocr_text, confidence = self._ocr_page(page)
                result.append(
                    TextSegment("PAGE", f"第 {index} 页（OCR）", ocr_text or text, confidence)
                )
        finally:
            document.close()
        return result

    def _ocr_page(self, page: fitz.Page) -> tuple[str, float | None]:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        data = pytesseract.image_to_data(
            image, lang="chi_sim+eng", config="--psm 6", output_type=pytesseract.Output.DICT
        )
        words: list[str] = []
        confidence_values: list[float] = []
        for text, confidence in zip(data["text"], data["conf"], strict=True):
            value = str(text).strip()
            try:
                score = float(confidence)
            except (TypeError, ValueError):
                score = -1
            if value:
                words.append(value)
            if score >= 0:
                confidence_values.append(score)
        average = (
            round(sum(confidence_values) / len(confidence_values), 2)
            if confidence_values
            else None
        )
        return " ".join(words), average

    def _read_xlsx(self, content: bytes) -> list[TextSegment]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        result: list[TextSegment] = []
        try:
            for sheet in workbook.worksheets:
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = ["" if value is None else str(value).strip() for value in row]
                    if any(values):
                        result.append(TextSegment("SHEET_ROW", f"{sheet.title}!{row_index}", " | ".join(values)))
        finally:
            workbook.close()
        return result

    def _read_csv(self, content: bytes) -> list[TextSegment]:
        text = self._decode(content)
        return [
            TextSegment("SHEET_ROW", f"行 {index}", " | ".join(row))
            for index, row in enumerate(csv.reader(io.StringIO(text)), start=1)
            if any(cell.strip() for cell in row)
        ]

    def _read_text(self, content: bytes) -> list[TextSegment]:
        text = self._decode(content)
        return [
            TextSegment("PARAGRAPH", f"段落 {index}", line.strip())
            for index, line in enumerate(text.splitlines(), start=1)
            if line.strip()
        ]

    def _read_legacy_doc(self, content: bytes) -> list[TextSegment]:
        with tempfile.TemporaryDirectory(prefix="czghagent-doc-") as directory:
            directory_path = Path(directory)
            source = directory_path / "source.doc"
            source.write_bytes(content)
            for converter in self._legacy_converters():
                try:
                    converted = self._convert_legacy_doc(converter, source, directory_path)
                    segments = self._read_text(converted) if converted else []
                    if segments:
                        return segments
                except (OSError, subprocess.TimeoutExpired, UnsupportedDocumentError):
                    continue
        raise UnsupportedDocumentError("旧版 DOC 文件转换失败，请另存为 DOCX 后重试")

    def _legacy_converters(self) -> list[str]:
        converters: list[str] = []
        for name in ("soffice", "libreoffice", "antiword", "textutil"):
            path = shutil.which(name)
            if path and path not in converters:
                converters.append(path)
        if not converters:
            raise UnsupportedDocumentError("运行环境未安装 DOC 转换器，请上传 DOCX 文件")
        return converters

    def _convert_legacy_doc(self, converter: str, source: Path, directory: Path) -> bytes | None:
        converter_name = Path(converter).name.lower()
        output_file = source.with_suffix(".txt")
        if converter_name in {"soffice", "libreoffice"}:
            profile = (directory / "libreoffice-profile").as_uri()
            command = [
                converter,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                str(directory),
                str(source),
            ]
        elif converter_name == "textutil":
            command = [converter, "-convert", "txt", "-stdout", str(source)]
        else:
            command = [converter, str(source)]
        environment = os.environ.copy()
        environment["XDG_CACHE_HOME"] = str(directory / "cache")
        completed = subprocess.run(
            command, capture_output=True, timeout=60, check=False, env=environment
        )
        if completed.returncode != 0:
            return None
        if converter_name in {"soffice", "libreoffice"}:
            return output_file.read_bytes() if output_file.exists() else None
        return completed.stdout or None

    def _decode(self, content: bytes) -> str:
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise UnsupportedDocumentError("文本编码无法识别")
