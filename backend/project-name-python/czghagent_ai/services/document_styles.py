# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-08-03
from collections.abc import Iterable
from typing import Any, Literal

from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

A4_TABLE_WIDTH_DXA = 9072
INK = "202124"
MUTED = "5F6368"
TABLE_HEADER = "E8EEF5"
LayoutProfile = Literal["standard", "compact", "condensed", "dense", "spacious"]


def configure_document(
    document: DocumentObject, profile: LayoutProfile = "standard"
) -> None:
    vertical_margin = {
        "standard": 25, "compact": 20, "condensed": 19, "dense": 18, "spacious": 27
    }[profile]
    horizontal_margin = {
        "standard": 25, "compact": 25, "condensed": 22, "dense": 20, "spacious": 25
    }[profile]
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(vertical_margin)
        section.bottom_margin = Mm(vertical_margin)
        section.left_margin = Mm(horizontal_margin)
        section.right_margin = Mm(horizontal_margin)
        section.header_distance = Mm(12.5)
        section.footer_distance = Mm(12.5)
        section.different_first_page_header_footer = True
    _configure_styles(document, profile)
    _configure_header_footer(document)
    _enable_field_updates(document)


def _configure_styles(document: DocumentObject, profile: LayoutProfile) -> None:
    body_font_size = {
        "standard": 12, "compact": 12, "condensed": 11, "dense": 10.5, "spacious": 12
    }[profile]
    normal_line_spacing = {
        "standard": 1.5, "compact": 1.35, "condensed": 1.25, "dense": 1.15,
        "spacious": 1.65,
    }[profile]
    normal_space_after = {
        "standard": 6, "compact": 4, "condensed": 3, "dense": 2, "spacious": 7
    }[profile]
    heading_spacing_scale = {
        "standard": 1.0, "compact": 0.8, "condensed": 0.7, "dense": 0.6,
        "spacious": 1.1,
    }[profile]
    heading_line_spacing = {
        "standard": 1.25, "compact": 1.15, "condensed": 1.1, "dense": 1.05,
        "spacious": 1.35,
    }[profile]
    normal = document.styles["Normal"]
    _set_style_font(normal, "宋体", body_font_size, INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(normal_space_after)
    normal.paragraph_format.line_spacing = normal_line_spacing
    heading_tokens = {
        1: ("黑体", 18, 18, 10),
        2: ("黑体", 16, 14, 8),
        3: ("黑体", 14, 10, 6),
        4: ("黑体", 13, 8, 4),
        5: ("楷体", 12, 6, 3),
        6: ("楷体", 12, 6, 3),
    }
    for level, (font, size, before, after) in heading_tokens.items():
        style = document.styles[f"Heading {level}"]
        _set_style_font(style, font, size, INK, bold=True)
        style.paragraph_format.space_before = Pt(before * heading_spacing_scale)
        style.paragraph_format.space_after = Pt(after * heading_spacing_scale)
        style.paragraph_format.line_spacing = heading_line_spacing
        style.paragraph_format.keep_with_next = True
    for level in range(1, 4):
        name = f"TOC {level}"
        style = (
            document.styles[name]
            if name in document.styles
            else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        )
        _set_style_font(
            style,
            "黑体" if level == 1 else "宋体",
            11.5 if level == 1 else 11,
            INK,
            bold=level == 1,
        )
        style.paragraph_format.left_indent = Mm((level - 1) * 7)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(5 if level == 1 else 3)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.tab_stops.add_tab_stop(
            Mm(157), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        _set_style_font(style, "宋体", body_font_size, INK)
        style.paragraph_format.left_indent = Mm(10)
        style.paragraph_format.first_line_indent = Mm(-5)
        style.paragraph_format.space_after = Pt(
            {"standard": 4, "compact": 3, "condensed": 2.5, "dense": 2, "spacious": 5}[profile]
        )
        style.paragraph_format.line_spacing = normal_line_spacing
    if "Tender Caption" not in document.styles:
        caption = document.styles.add_style("Tender Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = document.styles["Tender Caption"]
    _set_style_font(caption, "宋体", 10.5, INK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(
        {"standard": 10, "compact": 7, "condensed": 6, "dense": 5, "spacious": 11}[profile]
    )
    caption.paragraph_format.space_after = Pt(
        {"standard": 5, "compact": 3, "condensed": 2.5, "dense": 2, "spacious": 6}[profile]
    )
    caption.paragraph_format.keep_with_next = True
    if "Tender Table Note" not in document.styles:
        note = document.styles.add_style("Tender Table Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = document.styles["Tender Table Note"]
    _set_style_font(note, "宋体", 9.5, MUTED)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before = Pt(
        {"standard": 4, "compact": 3, "condensed": 2.5, "dense": 2, "spacious": 5}[profile]
    )
    note.paragraph_format.space_after = Pt(
        {"standard": 9, "compact": 6, "condensed": 5, "dense": 4, "spacious": 10}[profile]
    )


def _configure_header_footer(document: DocumentObject) -> None:
    for section in document.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.space_after = Pt(0)
        run = header.add_run("技术投标文件")
        set_run_font(run, "宋体", 9, MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = footer.add_run("第 ")
        set_run_font(prefix, "宋体", 9, MUTED)
        _append_field(footer, "PAGE")
        suffix = footer.add_run(" 页")
        set_run_font(suffix, "宋体", 9, MUTED)


def add_body_paragraph(
    document: DocumentObject, text: str, profile: LayoutProfile = "standard"
) -> Paragraph:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Mm(8.5)
    run = paragraph.add_run(text)
    set_run_font(
        run, "宋体",
        {
            "standard": 12, "compact": 12, "condensed": 11, "dense": 10.5,
            "spacious": 12,
        }[profile],
        INK,
    )
    return paragraph


def add_list(
    document: DocumentObject, items: Iterable[str], ordered: bool,
    profile: LayoutProfile = "standard",
) -> None:
    style = "List Number" if ordered else "List Bullet"
    for item in items:
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Mm(10)
        paragraph.paragraph_format.first_line_indent = Mm(-5)
        run = paragraph.add_run(item)
        set_run_font(
            run, "宋体",
            {
                "standard": 12, "compact": 12, "condensed": 11, "dense": 10.5,
                "spacious": 12,
            }[profile],
            INK,
        )


def add_data_table(
    document: DocumentObject, header: list[str], rows: list[list[str]],
    caption: str = "", note: str = "", profile: LayoutProfile = "standard"
) -> None:
    column_count = max([len(header), *(len(row) for row in rows)], default=0)
    if column_count == 0:
        return
    if caption:
        document.add_paragraph(caption, style="Tender Caption")
    normalized_header = header + [""] * (column_count - len(header))
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    has_header = any(normalized_header)
    table = document.add_table(rows=1 if has_header else 0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = _content_widths(normalized_header, normalized_rows)
    _set_table_geometry(table, widths, 120)
    if has_header:
        _fill_row(table.rows[0].cells, normalized_header, True, profile)
        _repeat_header(table.rows[0]._tr)
        _keep_row_together(table.rows[0]._tr)
    for row_values in normalized_rows:
        row = table.add_row()
        _fill_row(row.cells, row_values, False, profile)
        _keep_row_together(row._tr)
    _set_table_geometry(table, widths, 120)
    if note:
        document.add_paragraph(note, style="Tender Table Note")
    else:
        document.add_paragraph().paragraph_format.space_after = Pt(5)


def set_run_font(
    run: Run, name: str, size: float, color: str = INK, bold: bool | None = None
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)
    _remove_theme_fonts(fonts)


def _set_style_font(style: object, name: str, size: float, color: str, bold: bool = False) -> None:
    font = style.font  # type: ignore[attr-defined]
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    font.italic = False
    font.underline = False
    font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()  # type: ignore[attr-defined]
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)
    _remove_theme_fonts(fonts)


def _remove_theme_fonts(fonts: Any) -> None:
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        attribute = qn(f"w:{name}")
        if attribute in fonts.attrib:
            del fonts.attrib[attribute]


def _content_widths(header: list[str], rows: list[list[str]]) -> list[int]:
    columns = len(header) if header else len(rows[0])
    weights = []
    for index in range(columns):
        values = [header[index], *(row[index] for row in rows)]
        longest = max((_display_width(value) for value in values), default=1)
        weights.append(max(5, min(36, longest)))
    total = sum(weights)
    widths = [max(720, round(A4_TABLE_WIDTH_DXA * weight / total)) for weight in weights]
    widths[-1] += A4_TABLE_WIDTH_DXA - sum(widths)
    return widths


def _display_width(value: str) -> int:
    return sum(2 if ord(character) > 127 else 1 for character in value)


def _set_table_geometry(table: Table, widths: list[int], indent: int) -> None:
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:w"), str(indent))
    table_indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths, strict=True):
            _set_cell_width(cell, value)


def _fill_row(
    cells: tuple[_Cell, ...], values: list[str], header: bool, profile: LayoutProfile
) -> None:
    for cell, value in zip(cells, values, strict=True):
        _set_cell(cell, value, header, profile)
        if header:
            _shade_cell(cell, TABLE_HEADER)


def _set_cell(cell: _Cell, text: str, header: bool, profile: LayoutProfile) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = {
        "standard": 1.25, "compact": 1.1, "condensed": 1.05, "dense": 1.0,
        "spacious": 1.35,
    }[profile]
    run = paragraph.add_run(text)
    table_font_size = {
        "standard": 10.5, "compact": 10.5, "condensed": 10, "dense": 9.5,
        "spacious": 10.5,
    }[profile]
    set_run_font(run, "黑体" if header else "宋体", table_font_size, INK, bold=header)
    vertical = {
        "standard": 100, "compact": 60, "condensed": 45, "dense": 30, "spacious": 130
    }[profile]
    horizontal = {
        "standard": 120, "compact": 90, "condensed": 75, "dense": 60, "spacious": 140
    }[profile]
    _set_cell_margins(cell, vertical, horizontal, vertical, horizontal)


def _set_cell_width(cell: _Cell, value: int) -> None:
    width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        cell._tc.get_or_add_tcPr().append(width)
    width.set(qn("w:w"), str(value))
    width.set(qn("w:type"), "dxa")


def _set_cell_margins(cell: _Cell, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _shade_cell(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_header(row: object) -> None:
    properties = row.get_or_add_trPr()  # type: ignore[attr-defined]
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _keep_row_together(row: object) -> None:
    properties = row.get_or_add_trPr()  # type: ignore[attr-defined]
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _append_field(paragraph: Paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, display, end])
    set_run_font(run, "宋体", 9 if instruction == "PAGE" else 11, MUTED)


def _enable_field_updates(document: DocumentObject) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_page_break_section(document: DocumentObject) -> None:
    document.add_section(WD_SECTION.NEW_PAGE)
