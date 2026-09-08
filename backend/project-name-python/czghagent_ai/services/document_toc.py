# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-08-03
import hashlib
from collections.abc import Sequence

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

from czghagent_ai.document_models import DocumentOutlineNode
from czghagent_ai.services.document_styles import INK, MUTED, set_run_font
from czghagent_ai.services.outline_numbering import build_outline_labels, order_outline


def add_reference_toc(
    document: DocumentObject,
    outline: Sequence[DocumentOutlineNode],
    page_numbers: Sequence[int] | None = None,
) -> None:
    """Add an updateable Word TOC with cached clickable entries and page references."""
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("目录"), "黑体", 20, INK, bold=True)
    sorted_outline = order_outline(outline)
    labels = build_outline_labels(outline)
    numbers = list(page_numbers) if page_numbers is not None else [0] * len(sorted_outline)
    if len(numbers) != len(sorted_outline):
        raise ValueError("目录页码数量必须与目录节点数量一致")
    last_index = len(sorted_outline) - 1
    for index, (node, page_number) in enumerate(
        zip(sorted_outline, numbers, strict=True)
    ):
        _add_toc_entry(
            document,
            node,
            labels[node.id],
            page_number,
            field_start=index == 0,
            field_end=index == last_index,
        )


def heading_bookmark_name(node_id: str) -> str:
    digest = hashlib.sha1(node_id.encode("utf-8"), usedforsecurity=False).hexdigest()
    return f"_TenderHeading_{digest[:24]}"


def add_heading_bookmark(
    paragraph: Paragraph, bookmark_name: str, bookmark_id: int
) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    runs = paragraph.runs
    if runs:
        runs[0]._r.addprevious(start)
        runs[-1]._r.addnext(end)
        return
    paragraph._p.extend([start, end])


def _add_toc_entry(
    document: DocumentObject,
    node: DocumentOutlineNode,
    title: str,
    page_number: int,
    *,
    field_start: bool,
    field_end: bool,
) -> None:
    paragraph = document.add_paragraph(style=f"TOC {node.level}")
    paragraph.paragraph_format.left_indent = Mm((node.level - 1) * 7)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5 if node.level == 1 else 3)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Mm(157), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    if field_start:
        _append_toc_field_start(paragraph)
    bookmark = heading_bookmark_name(node.id)
    title_run = paragraph.add_run(title)
    set_run_font(
        title_run,
        "黑体" if node.level == 1 else "宋体",
        11.5 if node.level == 1 else 11,
        INK,
        bold=node.level == 1,
    )
    _wrap_run_in_hyperlink(paragraph, title_run, bookmark)
    tab_run = paragraph.add_run("\t")
    set_run_font(tab_run, "宋体", 11, MUTED)
    _append_page_reference(
        paragraph, bookmark, str(page_number if page_number > 0 else "0000")
    )
    if field_end:
        end_run = paragraph.add_run()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run._r.append(end)


def _append_toc_field_start(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instruction, separate])


def _wrap_run_in_hyperlink(
    paragraph: Paragraph, run: object, bookmark_name: str
) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name)
    hyperlink.set(qn("w:history"), "1")
    run_element = run._r  # type: ignore[attr-defined]
    paragraph._p.remove(run_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


def _append_page_reference(
    paragraph: Paragraph, bookmark_name: str, fallback: str
) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" PAGEREF {bookmark_name} \\h "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, display, end])
    set_run_font(run, "宋体", 11, MUTED)
