# GENERATED_BY_AI
# MODEL: claude-opus-5
# DATE: 2026-09-15
"""招标文件读取：每种格式都要落成带来源定位的文本片段，读不出来时说清楚原因。

这一层决定解读看到的是什么——读漏一张表、把 GB18030 读成乱码、OCR 结果被丢掉，
后面的模型再好也只能对着残缺的原文作答，而且不会报错。
"""

import io
import shutil
import subprocess
from collections.abc import Callable
from pathlib import Path

import fitz  # type: ignore[import-untyped]
import pytesseract  # type: ignore[import-untyped]
import pytest
from docx import Document
from openpyxl import Workbook

from czghagent_ai.services.file_readers import FileReaders, UnsupportedDocumentError


@pytest.fixture()
def readers() -> FileReaders:
    return FileReaders()


# ── 入口 ──────────────────────────────────────────────────────────────


def test_rejects_an_unsupported_format_by_name(readers: FileReaders) -> None:
    with pytest.raises(UnsupportedDocumentError, match=r"\.pptx"):
        readers.read("汇报.pptx", b"whatever")


def test_rejects_a_file_without_an_extension(readers: FileReaders) -> None:
    with pytest.raises(UnsupportedDocumentError, match="无扩展名"):
        readers.read("招标文件", b"whatever")


def test_matches_the_extension_case_insensitively(readers: FileReaders) -> None:
    segments = readers.read("招标文件.TXT", "项目名称：智慧园区".encode())

    assert [segment.text for segment in segments] == ["项目名称：智慧园区"]


def test_a_file_with_nothing_readable_is_an_error_not_an_empty_result(readers: FileReaders) -> None:
    with pytest.raises(UnsupportedDocumentError, match="没有可读取的文字内容"):
        readers.read("空白.txt", b"  \n\t\n  ")


# ── 纯文本与编码 ───────────────────────────────────────────────────────


def test_text_lines_keep_their_original_line_numbers_and_skip_blanks(readers: FileReaders) -> None:
    content = "﻿第一行\n\n  第三行  \n".encode()

    segments = readers.read("说明.md", content)

    assert [(s.locator_type, s.locator, s.text) for s in segments] == [
        ("PARAGRAPH", "段落 1", "第一行"),
        ("PARAGRAPH", "段落 3", "第三行"),
    ]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "gb18030", "utf-16"])
def test_reads_the_encodings_tender_documents_actually_arrive_in(readers: FileReaders, encoding: str) -> None:
    segments = readers.read("招标公告.txt", "采购预算：壹佰万元整".encode(encoding))

    assert segments[0].text == "采购预算：壹佰万元整"


def test_bytes_no_encoding_can_read_are_reported_as_such(readers: FileReaders) -> None:
    with pytest.raises(UnsupportedDocumentError, match="文本编码无法识别"):
        readers.read("损坏.txt", b"\xff")


def test_csv_rows_are_joined_cell_by_cell_and_blank_rows_skipped(readers: FileReaders) -> None:
    content = "序号,评分项,分值\n,,\n1,技术方案,30\n".encode("gb18030")

    segments = readers.read("评分表.csv", content)

    assert [(s.locator_type, s.locator, s.text) for s in segments] == [
        ("SHEET_ROW", "行 1", "序号 | 评分项 | 分值"),
        ("SHEET_ROW", "行 3", "1 | 技术方案 | 30"),
    ]


# ── Office ───────────────────────────────────────────────────────────


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("第一章 项目概况", level=1)
    document.add_paragraph("本项目为智慧园区建设。")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "评分项"
    table.cell(0, 1).text = "分值"
    table.cell(1, 0).text = "技术\n方案"
    table.cell(1, 1).text = "30"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_headings_paragraphs_and_table_rows_all_come_through(readers: FileReaders) -> None:
    segments = readers.read("招标文件.docx", _docx_bytes())

    by_type = [(s.locator_type, s.text) for s in segments]
    assert ("HEADING", "第一章 项目概况") in by_type
    assert ("PARAGRAPH", "本项目为智慧园区建设。") in by_type
    assert ("SHEET_ROW", "评分项 | 分值") in by_type
    assert ("SHEET_ROW", "技术 方案 | 30") in by_type, "单元格里的换行要压成空格，否则一行被拆成两段"
    rows = [s.locator for s in segments if s.locator_type == "SHEET_ROW"]
    assert rows == ["表 1 行 1", "表 1 行 2"]


def test_xlsx_rows_are_located_by_sheet_and_row_and_empty_rows_skipped(readers: FileReaders) -> None:
    workbook = Workbook()
    # 不用 workbook.active：它的类型是「工作表 | 图表页 | None」，图表页没有 append。
    workbook.remove(workbook.worksheets[0])
    first = workbook.create_sheet("评分标准")
    first.append(["评分项", "分值", None])
    first.append([None, None, None])
    first.append(["商务", 20, "含资质"])
    second = workbook.create_sheet("报价")
    second.append(["总价", 1000000])
    buffer = io.BytesIO()
    workbook.save(buffer)

    segments = readers.read("评分表.xlsx", buffer.getvalue())

    assert [(s.locator, s.text) for s in segments] == [
        ("评分标准!1", "评分项 | 分值 | "),
        ("评分标准!3", "商务 | 20 | 含资质"),
        ("报价!1", "总价 | 1000000"),
    ]


# ── PDF 与 OCR ────────────────────────────────────────────────────────


def _pdf_bytes(*page_texts: str) -> bytes:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = bytes(document.tobytes())
    document.close()
    return data


LONG_TEXT = "Section 1 Project overview and the scope of procurement for this tender."


def test_pdf_pages_with_real_text_are_read_directly_without_ocr(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(FileReaders, "_ocr_page", lambda self, page: pytest.fail("有文字层的页不该走 OCR"))

    segments = readers.read("招标文件.pdf", _pdf_bytes(LONG_TEXT))

    assert [(s.locator, s.ocr_confidence) for s in segments] == [("第 1 页", None)]
    assert "Project overview" in segments[0].text


def test_scanned_pages_fall_back_to_ocr_and_keep_the_confidence(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(FileReaders, "_ocr_page", lambda self, page: ("扫描件识别出的文字", 88.5))

    segments = readers.read("扫描件.pdf", _pdf_bytes(LONG_TEXT, ""))

    assert [(s.locator, s.text, s.ocr_confidence) for s in segments] == [
        ("第 1 页", segments[0].text, None),
        ("第 2 页（OCR）", "扫描件识别出的文字", 88.5),
    ]


def test_when_ocr_finds_nothing_the_little_text_layer_there_is_kept(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(FileReaders, "_ocr_page", lambda self, page: ("", None))

    segments = readers.read("封面.pdf", _pdf_bytes("Cover page"))

    assert [(s.locator, s.text) for s in segments] == [("第 1 页（OCR）", "Cover page")]


def test_ocr_confidence_averages_only_real_scores_and_drops_empty_words(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(pytesseract, "image_to_data", lambda *args, **kwargs: {
        "text": ["项目", "", "预算", "  "],
        "conf": ["90", "-1", "80.5", "not-a-number"],
    })
    document = fitz.open(stream=_pdf_bytes(""), filetype="pdf")
    try:
        text, confidence = readers._ocr_page(document[0])
    finally:
        document.close()

    assert text == "项目 预算"
    assert confidence == 85.25


# ── 旧版 DOC ─────────────────────────────────────────────────────────


def _which_only(*installed: str) -> Callable[[str], str | None]:
    """只「装了」这几个转换器的 shutil.which 替身。"""
    return lambda name: f"/usr/bin/{name}" if name in installed else None


def test_legacy_doc_without_any_converter_asks_for_docx(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(shutil, "which", lambda name: None)

    with pytest.raises(UnsupportedDocumentError, match="未安装 DOC 转换器"):
        readers.read("老文件.doc", b"\xd0\xcf\x11\xe0")


def test_legacy_doc_falls_through_a_failing_converter_to_the_next(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(shutil, "which", _which_only("antiword", "textutil"))
    tried: list[str] = []

    def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        tool = Path(command[0]).name
        tried.append(tool)
        if tool == "antiword":
            raise subprocess.TimeoutExpired(command, 60)
        return subprocess.CompletedProcess(command, 0, stdout="投标人须知\n".encode(), stderr=b"")

    monkeypatch.setattr(subprocess, "run", fake_run)

    segments = readers.read("老文件.doc", b"\xd0\xcf\x11\xe0")

    assert tried == ["antiword", "textutil"]
    assert [s.text for s in segments] == ["投标人须知"]


def test_libreoffice_output_is_read_from_the_converted_file(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/soffice" if name == "soffice" else None)

    def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        source = Path(command[-1])
        source.with_suffix(".txt").write_bytes("第一章 总则\n".encode())
        return subprocess.CompletedProcess(command, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(subprocess, "run", fake_run)

    assert [s.text for s in readers.read("老文件.doc", b"\xd0\xcf\x11\xe0")] == ["第一章 总则"]


def test_legacy_doc_that_no_converter_can_read_asks_to_resave_as_docx(
    readers: FileReaders, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(shutil, "which", _which_only("antiword", "textutil"))
    monkeypatch.setattr(
        subprocess, "run",
        lambda command, **kwargs: subprocess.CompletedProcess(command, 1, stdout=b"", stderr=b"boom"),
    )

    with pytest.raises(UnsupportedDocumentError, match="另存为 DOCX"):
        readers.read("老文件.doc", b"\xd0\xcf\x11\xe0")
