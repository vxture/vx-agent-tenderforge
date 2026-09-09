# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-08-02
import io
import time

from docx import Document
from fastapi.testclient import TestClient

from czghagent_ai.app import app
from czghagent_ai.config import settings
from czghagent_ai.services.document_parser import DocumentParserService

client = TestClient(app)


# 组织规范 025 §3 的必填身份字段。守的是**字段名**：跨产品聚合按名字取值，
# 名字漂了聚合端读到空，而空值和「这个服务没部署」在那边长得一模一样。
_REQUIRED_IDENTITY = {"status", "service", "version", "gitSha", "stage", "buildTime", "time"}


def test_health_returns_the_full_identity_block() -> None:
    response = client.get("/health")
    assert response.status_code == 200
    body = response.json()
    assert _REQUIRED_IDENTITY <= body.keys()
    assert body["status"] == "ok"          # 规范定的是字面量 ok，不是 UP
    assert body["service"] == "tenderforge-ai"
    assert body["product"] == "tenderforge"


def test_health_falls_back_honestly_when_build_info_was_not_injected(
    monkeypatch: object,
) -> None:
    # 规范 §6 把编造版本号列为禁止项。没注入就如实报 unknown。
    for key in ("APP_VERSION", "GIT_SHA", "BUILD_TIME", "DEPLOY_STAGE"):
        monkeypatch.delenv(key, raising=False)  # type: ignore[attr-defined]
    body = client.get("/health").json()
    assert body["version"] == "dev"
    assert body["gitSha"] == "unknown"
    assert body["buildTime"] == "unknown"
    assert body["stage"] == "local"


def test_health_strips_the_image_tag_prefix_from_git_sha(monkeypatch: object) -> None:
    # `sha-` 是镜像 tag 的形态，不是数据——带着它聚合端拿到的值没法直接 git show。
    monkeypatch.setenv("GIT_SHA", "sha-763c71c")  # type: ignore[attr-defined]
    assert client.get("/health").json()["gitSha"] == "763c71c"


def test_health_time_is_taken_fresh_so_it_proves_the_clock() -> None:
    first = client.get("/health").json()["time"]
    time.sleep(1.1)
    second = client.get("/health").json()["time"]
    assert first != second


def test_parse_requires_internal_token() -> None:
    response = client.post(
        "/internal/parse",
        files={"file": ("招标文件.txt", "评分办法：技术方案满分30分。".encode(), "text/plain")},
    )
    assert response.status_code == 401


def test_parse_text_returns_source_located_evidence() -> None:
    response = client.post(
        "/internal/parse",
        headers={"X-Internal-Token": settings.internal_token},
        files={
            "file": (
                "招标文件.txt",
                "评分办法：技术方案满分30分。\n废标条款：未盖章作无效投标处理。".encode(),
                "text/plain",
            )
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["category"] == "TENDER_DOCUMENT"
    assert len(body["evidence"]) == 2
    assert body["evidence"][0]["locator"] == "段落 1"
    assert body["facts"] == []


def test_parse_docx_preserves_heading_table_and_paragraph_order() -> None:
    document = Document()
    document.add_heading("技术方案响应要求", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "评分项"
    table.cell(0, 1).text = "20分"
    document.add_paragraph("表后说明")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = DocumentParserService().parse("招标文件.docx", buffer.getvalue())

    assert len(parsed.evidence) == 3
    assert parsed.evidence[0].locator_type == "HEADING"
    assert parsed.evidence[1].locator_type == "SHEET_ROW"
    assert parsed.evidence[2].locator_type == "PARAGRAPH"
    assert [item.excerpt for item in parsed.evidence] == [
        "技术方案响应要求",
        "评分项 | 20分",
        "表后说明",
    ]
