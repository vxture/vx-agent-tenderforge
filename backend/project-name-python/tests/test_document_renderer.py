# GENERATED_BY_AI
# MODEL: gpt-5
# DATE: 2026-08-03
import io
import zipfile

from docx import Document

from czghagent_ai.document_models import DocumentOutlineNode, DocumentQaResult, DocumentRenderRequest
from czghagent_ai.services.document_html import parse_tender_html
from czghagent_ai.services.document_renderer import DocumentRenderer, _balance_cover_title


class PassingQualityAssurance:
    def resolve_outline_pages(
        self, docx: bytes, outline: list[DocumentOutlineNode]
    ) -> list[int]:
        assert docx.startswith(b"PK")
        return [3 + index for index, _ in enumerate(outline)]

    def inspect(
        self, docx: bytes, target_pages: int, forbidden_terms: list[str]
    ) -> DocumentQaResult:
        assert docx.startswith(b"PK")
        assert target_pages == 20
        assert forbidden_terms == ["16GB"]
        return DocumentQaResult(
            status="PASSED", actual_pages=20,
            summary="目标20页，实测20页，偏差+0页",
        )


class AdaptiveQualityAssurance(PassingQualityAssurance):
    def __init__(self) -> None:
        self.inspect_calls = 0
        self.resolve_calls = 0

    def resolve_outline_pages(
        self, docx: bytes, outline: list[DocumentOutlineNode]
    ) -> list[int]:
        self.resolve_calls += 1
        return super().resolve_outline_pages(docx, outline)

    def inspect(
        self, docx: bytes, target_pages: int, forbidden_terms: list[str]
    ) -> DocumentQaResult:
        self.inspect_calls += 1
        if self.inspect_calls == 1:
            return DocumentQaResult(
                status="FAILED", actual_pages=23,
                summary="目标20页，实测23页，偏差+3页",
            )
        return DocumentQaResult(
            status="PASSED", actual_pages=20,
            summary="目标20页，实测20页，偏差+0页",
        )


class DenseQualityAssurance(AdaptiveQualityAssurance):
    def inspect(
        self, docx: bytes, target_pages: int, forbidden_terms: list[str]
    ) -> DocumentQaResult:
        self.inspect_calls += 1
        actual_pages = {1: 30, 2: 26, 3: 24}.get(self.inspect_calls, 22)
        return DocumentQaResult(
            status="PASSED" if self.inspect_calls == 4 else "FAILED",
            actual_pages=actual_pages,
            summary=f"目标20页，实测{actual_pages}页",
        )


def test_renderer_numbers_outline_tables_and_preserves_privacy() -> None:
    request = DocumentRenderRequest.model_validate({
        "bidId": "bid-1",
        "title": "智能决策项目",
        "biddingMode": "BLIND",
        "targetPages": 20,
        "outline": [
            {"id": "n1", "parentId": None, "level": 1, "title": "总体方案", "sortOrder": 0},
            {"id": "n2", "parentId": "n1", "level": 2, "title": "技术路线", "sortOrder": 1},
            {"id": "n3", "parentId": "n2", "level": 3, "title": "模型训练", "sortOrder": 2},
        ],
        "chapters": [{
            "id": "c1", "outlineNodeId": "n3", "title": "模型训练",
            "content": (
                "<h4>一、组织措施</h4><p>（一）执行要求</p>"
                "<p>采用分层训练与持续评估机制。</p>"
                '<p data-table-title="true">表 1 资源配置</p>'
                "<table><thead><tr><th>资源</th><th>配置</th></tr></thead>"
                "<tbody><tr><td>前置代理服务器</td><td>8核/64GB/300GB</td></tr></tbody></table>"
                '<p data-table-note="true">注：容量为最低配置要求。</p>'
                '<div data-flow-diagram="legacy"><p>历史流程图</p></div>'
            ),
        }],
        "forbiddenTerms": ["16GB"],
    })

    renderer = DocumentRenderer(PassingQualityAssurance())  # type: ignore[arg-type]
    content, qa = renderer.render(request)

    assert qa.status == "PASSED"
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "第一章 总体方案" in paragraphs
    assert "一、技术路线" in paragraphs
    assert "（一）模型训练" in paragraphs
    assert "采用分层训练与持续评估机制。" in paragraphs
    assert "1. 组织措施" in paragraphs
    assert "（1）执行要求" in paragraphs
    assert not any("一、组织措施" in paragraph for paragraph in paragraphs)
    assert "第一章 总体方案\t3" in paragraphs
    assert "一、技术路线\t4" in paragraphs
    assert "（一）模型训练\t5" in paragraphs
    assert "表 1-1 资源配置" in paragraphs
    assert "注：容量为最低配置要求。" in paragraphs
    assert "历史流程图" not in paragraphs
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == "8核/64GB/300GB"
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        document_xml = archive.read("word/document.xml")
        settings_xml = archive.read("word/settings.xml")
        styles_xml = archive.read("word/styles.xml")
        assert b"rsid" not in document_xml
        assert document_xml.count(b"w:cantSplit") == 2
        assert b'TOC \\o "1-3" \\h \\z \\u' in document_xml
        assert document_xml.count(b" PAGEREF ") == 3
        assert document_xml.count(b"<w:hyperlink ") == 3
        assert document_xml.count(b"<w:bookmarkStart ") == 3
        assert b'w:updateFields w:val="true"' in settings_xml
        assert "0000" not in document_xml.decode("utf-8")
        assert "投标人" not in document_xml.decode("utf-8")
        assert "MS Gothic" not in document_xml.decode("utf-8")
        assert document_xml.decode("utf-8").count('w:eastAsia="黑体"') >= 3
        for heading_style in ("Heading1", "Heading2", "Heading3"):
            style = styles_xml.decode("utf-8").split(f'w:styleId="{heading_style}"', 1)[1]
            style = style.split("</w:style>", 1)[0]
            assert 'w:ascii="黑体"' in style
            assert 'w:hAnsi="黑体"' in style
            assert 'w:eastAsia="黑体"' in style
            assert "Theme=" not in style
        for toc_style in ("TOC1", "TOC2", "TOC3"):
            assert f'w:styleId="{toc_style}"' in styles_xml.decode("utf-8")


def test_renderer_reflows_page_only_qa_failure_with_compact_profile() -> None:
    request = DocumentRenderRequest.model_validate({
        "bidId": "bid-1",
        "title": "智能决策项目",
        "biddingMode": "BLIND",
        "targetPages": 20,
        "outline": [
            {"id": "n1", "parentId": None, "level": 1, "title": "总体方案", "sortOrder": 0},
            {"id": "n2", "parentId": "n1", "level": 2, "title": "技术路线", "sortOrder": 1},
            {"id": "n3", "parentId": "n2", "level": 3, "title": "模型训练", "sortOrder": 2},
        ],
        "chapters": [{
            "id": "c1", "outlineNodeId": "n3", "title": "模型训练",
            "content": "<p>采用分层训练与持续评估机制。</p>",
        }],
        "forbiddenTerms": ["16GB"],
    })
    quality_assurance = AdaptiveQualityAssurance()

    _, qa = DocumentRenderer(quality_assurance).render(request)  # type: ignore[arg-type]

    assert qa.status == "PASSED"
    assert qa.actual_pages == 20
    assert "已自动采用紧凑排版" in qa.summary
    assert quality_assurance.resolve_calls == 2
    assert quality_assurance.inspect_calls == 2


def test_renderer_uses_dense_profile_after_compact_profile_still_fails() -> None:
    request = DocumentRenderRequest.model_validate({
        "bidId": "bid-1",
        "title": "智能决策项目",
        "biddingMode": "BLIND",
        "targetPages": 20,
        "outline": [
            {"id": "n1", "parentId": None, "level": 1, "title": "总体方案", "sortOrder": 0},
            {"id": "n2", "parentId": "n1", "level": 2, "title": "技术路线", "sortOrder": 1},
            {"id": "n3", "parentId": "n2", "level": 3, "title": "模型训练", "sortOrder": 2},
        ],
        "chapters": [{
            "id": "c1", "outlineNodeId": "n3", "title": "模型训练",
            "content": "<p>采用分层训练与持续评估机制。</p>",
        }],
        "forbiddenTerms": ["16GB"],
    })
    quality_assurance = DenseQualityAssurance()

    _, qa = DocumentRenderer(quality_assurance).render(request)  # type: ignore[arg-type]

    assert qa.status == "PASSED"
    assert qa.actual_pages == 22
    assert "已自动采用密排排版" in qa.summary
    assert quality_assurance.resolve_calls == 4
    assert quality_assurance.inspect_calls == 4


def test_markdown_pipe_table_is_converted_with_caption_and_note() -> None:
    blocks = parse_tender_html(
        "<p>表1 系统对接范围</p>"
        "<p>| 系统 | 接口 |\n|---|---|\n| ERP | RESTful API |</p>"
        "<p>表注：同步频次按业务要求配置。</p>"
    )

    assert len(blocks) == 1
    assert blocks[0].kind == "table"
    assert blocks[0].caption == "表1 系统对接范围"
    assert blocks[0].header == ["系统", "接口"]
    assert blocks[0].rows == [["ERP", "RESTful API"]]
    assert blocks[0].note == "表注：同步频次按业务要求配置。"


def test_duplicate_legacy_table_titles_keep_the_first_semantic_title() -> None:
    blocks = parse_tender_html(
        '<p data-table-title="true">验收测试内容与通过准则表</p>'
        '<p data-table-title="true">验收测试与评审明细表</p>'
        "<table><thead><tr><th>测试类型</th><th>通过准则</th></tr></thead>"
        "<tbody><tr><td>功能测试</td><td>全部通过</td></tr></tbody></table>"
    )

    assert len(blocks) == 1
    assert blocks[0].caption == "验收测试内容与通过准则表"


def test_renderer_numbers_tables_continuously_within_each_top_level_chapter() -> None:
    request = DocumentRenderRequest.model_validate({
        "bidId": "bid-table-numbering",
        "title": "智能决策项目",
        "biddingMode": "BLIND",
        "targetPages": 20,
        "outline": [
            {"id": "r1", "parentId": None, "level": 1, "title": "总体方案", "sortOrder": 0},
            {"id": "s1", "parentId": "r1", "level": 2, "title": "技术路线", "sortOrder": 1},
            {"id": "l1", "parentId": "s1", "level": 3, "title": "资源组织", "sortOrder": 2},
            {"id": "l2", "parentId": "s1", "level": 3, "title": "运行保障", "sortOrder": 3},
            {"id": "r2", "parentId": None, "level": 1, "title": "验收方案", "sortOrder": 4},
            {"id": "s2", "parentId": "r2", "level": 2, "title": "验收组织", "sortOrder": 5},
            {"id": "l3", "parentId": "s2", "level": 3, "title": "验收方法", "sortOrder": 6},
        ],
        "chapters": [
            {
                "id": "c1", "outlineNodeId": "l1", "title": "资源组织",
                "content": (
                    '<p data-table-title="true">资源配置</p>'
                    '<table><tr><th>资源</th><th>配置</th></tr><tr><td>A</td><td>B</td></tr></table>'
                    '<p data-table-note="true"></p>'
                    '<p data-table-title="true">表 9-3 实施分工</p>'
                    '<table><tr><th>角色</th><th>职责</th></tr><tr><td>A</td><td>B</td></tr></table>'
                    '<p data-table-note="true"></p>'
                ),
            },
            {
                "id": "c2", "outlineNodeId": "l2", "title": "运行保障",
                "content": (
                    '<table><tr><th>事项</th><th>要求</th></tr><tr><td>A</td><td>B</td></tr></table>'
                ),
            },
            {
                "id": "c3", "outlineNodeId": "l3", "title": "验收方法",
                "content": (
                    '<p data-table-title="true">表1 验收准则</p>'
                    '<table><tr><th>事项</th><th>准则</th></tr><tr><td>A</td><td>B</td></tr></table>'
                ),
            },
        ],
        "forbiddenTerms": ["16GB"],
    })

    content, _ = DocumentRenderer(PassingQualityAssurance()).render(request)  # type: ignore[arg-type]
    paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]

    assert "表 1-1 资源配置" in paragraphs
    assert "表 1-2 实施分工" in paragraphs
    assert "表 1-3 运行保障明细表" in paragraphs
    assert "表 2-1 验收准则" in paragraphs
    assert not any("表 9-3" in paragraph or "表 2-1 表1" in paragraph for paragraph in paragraphs)


def test_collapsed_markdown_table_is_rendered_as_native_docx_table() -> None:
    markdown = (
        "表1 高性能推算服务运维内容 | 运维层面 | 服务内容 | 响应要求 | "
        "|----------|----------|----------| | 硬件运维 | 设备状态监控 | 7×24小时现场响应 | "
        "| 系统运维 | 补丁管理 | 每月一次例行维护 | | 应用运维 | 性能调优 | 按需提供 |"
    )
    request = DocumentRenderRequest.model_validate({
        "bidId": "bid-inline-table",
        "title": "智能决策项目",
        "biddingMode": "BLIND",
        "targetPages": 20,
        "outline": [
            {"id": "n1", "parentId": None, "level": 1, "title": "总体方案", "sortOrder": 0},
            {"id": "n2", "parentId": "n1", "level": 2, "title": "运维服务", "sortOrder": 1},
            {"id": "n3", "parentId": "n2", "level": 3, "title": "运维内容", "sortOrder": 2},
        ],
        "chapters": [{
            "id": "c1", "outlineNodeId": "n3", "title": "运维内容",
            "content": f"<p>{markdown}</p>",
        }],
        "forbiddenTerms": ["16GB"],
    })

    content, qa = DocumentRenderer(PassingQualityAssurance()).render(request)  # type: ignore[arg-type]
    document = Document(io.BytesIO(content))

    assert qa.status == "PASSED"
    assert "表 1-1 高性能推算服务运维内容" in [item.text for item in document.paragraphs]
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "运维层面"
    assert document.tables[0].cell(1, 0).text == "硬件运维"
    assert document.tables[0].cell(3, 2).text == "按需提供"


def test_collapsed_numbered_content_is_split_before_docx_rendering() -> None:
    blocks = parse_tender_html(
        "<p>1. 宣传部智能体设计 （1）建设定位 正文内容。"
        " 1）宣传文稿智能生成 功能说明。 2. 组织部智能体设计 正文内容。</p>"
    )

    assert [block.text for block in blocks] == [
        "1. 宣传部智能体设计",
        "（1）建设定位 正文内容。",
        "（1）宣传文稿智能生成 功能说明。",
        "2. 组织部智能体设计 正文内容。",
    ]


def test_chinese_ordinal_items_are_split_before_docx_rendering() -> None:
    blocks = parse_tender_html(
        "<p>建设目标明确而系统。第一，建设统一平台底座。"
        " 第二，建立长效运营机制。 第三，编制五年建设规划。</p>"
    )

    assert [block.text for block in blocks] == [
        "建设目标明确而系统。",
        "第一，建设统一平台底座。",
        "第二，建立长效运营机制。",
        "第三，编制五年建设规划。",
    ]


def test_cover_title_does_not_split_ascii_number_or_measure_word() -> None:
    balanced = _balance_cover_title(
        "金钼生产运营辅助智能决策项目-80页暗标质量验收-20260807"
    )

    assert "8\n0" not in balanced
    assert "80\n页" not in balanced
